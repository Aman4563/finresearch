"""Renderer: citation links, evidence appendix, tables, pack layout, gate naming, and safe PDF re-renders."""

from __future__ import annotations

import hashlib
import shutil
from decimal import Decimal
from pathlib import Path

import pytest
from pypdf import PdfReader

from finresearch.render.html import (
    EXPORT_DISCLAIMER,
    EXPORT_WATERMARK,
    ClaimView,
    evidence_appendix,
    link_citations,
    render_html,
)
from finresearch.render.pack import financial_pivot
from finresearch.render.pdf import find_chrome, html_to_pdf

FIX = Path(__file__).parent / "fixtures"


def test_citation_links_carry_status_and_missing_claims_are_red():
    claims = {1: ClaimView(1, "verified", "financials", "PAT ₹535.61 mn",
                           citations=[{"doc_title": "RHP", "page": 73, "lines": "10-10", "quote": "Profit 535.61"}]),
              2: ClaimView(2, "contradicted", "demand", "GMP ₹90")}  # fmt: skip
    out = link_citations("<p>PAT [C1], GMP [C2], ghost [C9]</p>", claims)
    assert 'class="cite verified" href="#claim-1"' in out and "RHP p73 L10-10: Profit 535.61" in out
    assert 'class="cite contradicted"' in out and 'class="cite missing"' in out


def test_evidence_appendix_shows_quotes_urls_and_corrections():
    html = evidence_appendix([
        ClaimView(5, "verified", "financials", "PAT", value="535.61", unit="INR mn", period="FY26",
                  citations=[{"doc_title": "RHP", "page": 73, "lines": "4966-4966", "quote": "Profit 535.61",
                              "quote_found": True}]),
        ClaimView(6, "needs_review", "demand", "GMP", corrects=4, note="stale",
                  citations=[{"url": "https://gmp.example/x", "accessed_at": "2026-09-28T13:37"}]),
    ])  # fmt: skip
    assert "id='claim-5'" in html and "535.61 INR mn (FY26)" in html and "✓ quote found" in html
    assert "https://gmp.example/x" in html and "corrects C4" in html and "verifier: stale" in html


def test_export_html_has_disclaimer_and_watermark_even_when_blocked():
    from datetime import datetime

    for ok in (True, False):
        out = render_html("# R\nText [C1].", {}, title="T", generated_at=datetime(2026, 9, 30), gate_ok=ok,
                          gate_blocking=["x"], gate_warnings=[])  # fmt: skip
        assert out.count(EXPORT_DISCLAIMER) == 2 and EXPORT_WATERMARK in out and "aria-hidden='true'" in out
        # the report body and banner are unchanged: the notice is added around them
        assert ("Publish gate passed" in out) is ok and ("NOT PUBLISHED" in out) is not ok


def test_financial_pivot_prefers_verified_and_drops_contradicted():
    claims = {
        1: ClaimView(
            1, "unverified", "financials", "", metric="revenue", value="100", unit="INR mn", period="FY25"
        ),
        2: ClaimView(
            2, "verified", "financials", "", metric="revenue", value="101", unit="INR mn", period="FY25"
        ),
        3: ClaimView(
            3, "contradicted", "financials", "", metric="revenue", value="999", unit="INR mn", period="FY26"
        ),
        4: ClaimView(4, "verified", "demand", "", metric="revenue", value="5", unit="INR mn", period="FY26"),
    }
    header, rows = financial_pivot(claims)
    assert header == ["metric", "FY25"] and rows == [["revenue (INR mn)", "101 [C2]"]]


@pytest.fixture
def run_with_report(env, tmp_path):
    from finresearch.db import session_scope
    from finresearch.db.models import AgentStep, Citation, Claim, Document, DocumentPage, ResearchRun
    from finresearch.ingest.documents import get_or_create_company

    txt = tmp_path / "text.txt"
    shutil.copy(FIX / "orient_rhp_pnl_layout.txt", txt)
    pdf_src = tmp_path / "rhp.pdf"
    pdf_src.write_bytes(b"%PDF-1.4 fake")
    lines = txt.read_text().split("\n")
    ln = next(i for i, x in enumerate(lines, 1) if "Profit / (Loss) for the year" in x)
    with session_scope() as s:
        co = get_or_create_company(
            s, "render-" + hashlib.sha1(str(tmp_path).encode()).hexdigest()[:8], "Render Co"
        )
        d = Document(company_id=co.id, kind="RHP", title="Render RHP", sha256=hashlib.sha256(str(tmp_path).encode())
                     .hexdigest(), local_path=str(pdf_src), text_path=str(txt), bytes=1, pages=1)  # fmt: skip
        s.add(d)
        s.flush()
        s.add(DocumentPage(document_id=d.id, page_no=73, line_start=1, line_end=len(lines), text="", char_count=0,
                           text_source="pdftotext"))  # fmt: skip
        run = ResearchRun(company_id=co.id, kind="ipo_report", status="done", manifest={})
        s.add(run)
        s.flush()
        ids = []
        for period, val in (("FY24", "400.69"), ("FY25", "533.21"), ("FY26", "535.61")):
            c = Claim(run_id=run.id, stream="financials", statement=f"PAT {period}", claim_type="numeric", metric="pat",
                      value=Decimal(val), unit="INR million", period=period, status="verified")  # fmt: skip
            s.add(c)
            s.flush()
            s.add(Citation(claim_id=c.id, document_id=d.id, page_no=73, line_start=ln, line_end=ln,
                           quote=lines[ln - 1].strip(), quote_found=True))  # fmt: skip
            ids.append(c.id)
        report = f"# Render Co report\n\n## 1. Verdict\nPAT rose to ₹535.61 mn [C{ids[2]}] from ₹400.69 mn [C{ids[0]}].\n"
        s.add(AgentStep(run_id=run.id, key="synthesis", stage="synthesis", role="synthesizer", status="done",
                        output={"report_markdown": report}))  # fmt: skip
        s.add(AgentStep(run_id=run.id, key="stream:financials", stage="stream", role="financials", status="done",
                        output={"section_markdown": f"## Financials\nPAT [C{ids[2]}]"}))  # fmt: skip
        return run.id, co.slug


def test_render_pack_layout_tables_and_published_names(run_with_report, env):
    from openpyxl import load_workbook

    from finresearch.render.pack import render_pack

    run_id, slug = run_with_report
    r = render_pack(run_id, pdf=False)
    assert r.gate_ok and r.path == env.reports_dir / slug / f"run-{run_id}"
    for name in ("01_Offer_Documents/RHP__Render_RHP.pdf", "02_Financial_Reports/financials.xlsx",
                 "02_Financial_Reports/financials_section.md", "06_Final_Report/report.md",
                 "06_Final_Report/report.html", "06_Final_Report/fact_check_log.md", "06_Final_Report/claims.xlsx",
                 "README.md"):  # fmt: skip
        assert (r.path / name).exists(), name
    assert list((r.path / "02_Financial_Reports" / "charts").glob("*.png")), "PAT has 3 periods -> one chart"
    wb = load_workbook(r.path / "02_Financial_Reports" / "financials.xlsx")
    ws = wb["Financials"]
    assert [c.value for c in ws[1]] == ["metric", "FY24", "FY25", "FY26"] and ws["D2"].value.startswith(
        "535.61 [C"
    )
    html = (r.path / "06_Final_Report" / "report.html").read_text()
    assert (
        "Publish gate passed" in html and "Evidence: claims cited" in html and "data:image/png;base64" in html
    )
    assert "PASSED" in (r.path / "README.md").read_text()
    # §D.8: every export carries the personal-use disclaimer (top and bottom) and the watermark
    readme = (r.path / "README.md").read_text()
    assert html.count(EXPORT_DISCLAIMER) == 2 and "class='watermark'" in html and EXPORT_WATERMARK in html
    assert EXPORT_DISCLAIMER in readme and "not a SEBI-registered" in EXPORT_DISCLAIMER


def test_blocked_report_is_rendered_as_not_published(run_with_report):
    from finresearch.db import session_scope
    from finresearch.db.models import AgentStep
    from finresearch.render.pack import render_pack

    run_id, _ = run_with_report
    with session_scope() as s:
        st = s.query(AgentStep).filter_by(run_id=run_id, key="synthesis").one()
        st.output = {"report_markdown": "# R\nPAT ₹535.61 mn [RHP L4966] and [C999999]."}
    r = render_pack(run_id, pdf=False)
    final = r.path / "06_Final_Report"
    assert (
        not r.gate_ok and not (final / "report.md").exists() and (final / "report_NOT_PUBLISHED.md").exists()
    )
    html = (final / "report_NOT_PUBLISHED.html").read_text()
    assert (
        "NOT PUBLISHED" in html
        and "raw document citations" in html
        and "BLOCKED" in (r.path / "README.md").read_text()
    )


@pytest.mark.skipif(find_chrome() is None, reason="Chrome not available")
def test_pdf_rerender_over_existing_file_stays_valid(tmp_path):
    """Regression: re-rendering over an existing PDF once produced two concatenated documents."""
    html = tmp_path / "r.html"
    html.write_text("<html><body><h1>Report</h1>" + "<p>line</p>" * 400 + "</body></html>")
    pdf = tmp_path / "r.pdf"
    n1 = html_to_pdf(html, pdf, title="t")
    size1 = pdf.stat().st_size
    n2 = html_to_pdf(html, pdf, title="t")
    reader = PdfReader(pdf, strict=True)
    assert n1 == n2 == len(reader.pages) > 1 and abs(pdf.stat().st_size - size1) < size1 * 0.1
    assert pdf.read_bytes().count(b"%%EOF") == 1
