"""Build the research folder pack for a run (same layout as the manual research packs).

<reports_dir>/<company>/run-<id>/
  README.md
  01_Offer_Documents/        RHP, DRHP, addenda, abridged prospectus, price-band ad, anchor letter, industry report
  02_Financial_Reports/      annual reports + statements, financials section, financials.xlsx/.csv, charts/
  03_News_Last_30_Days/      news section + source list
  04_Major_News_and_Events/  risks + major-history sections
  05_Stock_and_Valuation_Analysis/  valuation, demand, business sections
  06_Final_Report/           report.md/.html/.pdf, fact_check_log.md, claims.xlsx/.csv
"""

from __future__ import annotations

import base64
import csv
import io
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from sqlalchemy import select

from finresearch.config import get_settings
from finresearch.db import session_scope
from finresearch.db.models import AgentStep, Claim, Company, Document, ResearchRun
from finresearch.fincalc.dates import now_ist
from finresearch.render.html import EXPORT_DISCLAIMER, EXPORT_WATERMARK, ClaimView, render_html
from finresearch.render.pdf import PdfRenderError, html_to_pdf
from finresearch.verify.gate import check_report

FOLDERS = {
    "01": "01_Offer_Documents",
    "02": "02_Financial_Reports",
    "03": "03_News_Last_30_Days",
    "04": "04_Major_News_and_Events",
    "05": "05_Stock_and_Valuation_Analysis",
    "06": "06_Final_Report",
}
DOC_FOLDER = {"RHP": "01", "DRHP": "01", "ADDENDUM": "01", "ABRIDGED_PROSPECTUS": "01", "PRICE_BAND_AD": "01",
              "ANCHOR_ALLOCATION": "01", "INDUSTRY_REPORT": "01", "ANNUAL_REPORT": "02",
              "FINANCIAL_STATEMENTS": "02", "OTHER": "05"}  # fmt: skip
STREAM_FOLDER = {"financials": "02", "news30": "03", "risks": "04", "major": "04", "valuation": "05",
                 "demand": "05", "business": "05"}  # fmt: skip


@dataclass(frozen=True)
class PackLayout:
    """Where a research kind's documents and stream sections go in its research pack."""

    label: str
    folders: dict[str, str]
    doc_folder: dict[str, str]
    stream_folder: dict[str, str]
    financial_streams: tuple[str, ...]
    readme_rows: tuple[tuple[str, str], ...]
    disclaimer: str


IPO_LAYOUT = PackLayout(
    label="IPO research report", folders=FOLDERS, doc_folder=DOC_FOLDER, stream_folder=STREAM_FOLDER,
    financial_streams=("financials",),
    readme_rows=(("01", "RHP, DRHP, addenda, abridged prospectus, price-band ad, anchor letter, industry report"),
                 ("02", "Annual reports and statements, financials section, financials.xlsx/.csv, charts"),
                 ("03", "News section and every web source cited"),
                 ("04", "Risks, litigation and governance; history, sector and macro"),
                 ("05", "Valuation and peers, demand signals, business and industry"),
                 ("06", "Report (md/html/pdf), fact-check log, full claim ledger (xlsx/csv)")),
    disclaimer="Personal research, not SEBI-registered investment advice. Live figures are INTERIM; GMP is unofficial.",
)  # fmt: skip
STOCK_LAYOUT = PackLayout(
    label="stock research report",
    folders={"01": "01_Company_Filings", "02": "02_Financial_Reports", "03": "03_News_Last_30_Days",
             "04": "04_Governance_and_Events", "05": "05_Valuation_and_Price", "06": "06_Final_Report"},
    doc_folder={"ANNUAL_REPORT": "01", "FINANCIAL_STATEMENTS": "02", "OTHER": "01"},
    stream_folder={"stock_fundamentals": "02", "stock_news": "03", "stock_governance": "04", "stock_business": "05",
                   "stock_valuation": "05", "stock_technical": "05"},
    financial_streams=("stock_fundamentals",),
    readme_rows=(("01", "Annual reports and other company filings"),
                 ("02", "Results filings, fundamentals section, financials.xlsx/.csv, charts"),
                 ("03", "News section and every web source cited"),
                 ("04", "Governance, ownership and major events"),
                 ("05", "Business, valuation vs history and peers, price behaviour"),
                 ("06", "Report (md/html/pdf), fact-check log, full claim ledger (xlsx/csv)")),
    disclaimer="Personal research, not SEBI-registered investment advice. Prices and holdings are as of their dates.",
)  # fmt: skip
FUND_LAYOUT = PackLayout(
    label="mutual-fund research report",
    folders={"01": "01_Scheme_Documents", "02": "02_Performance_and_Risk", "03": "03_News",
             "04": "04_Management_and_AMC", "05": "05_Portfolio_and_Costs", "06": "06_Final_Report"},
    doc_folder={"OTHER": "01"},
    stream_folder={"fund_performance": "02", "fund_risk": "02", "fund_news": "03", "fund_manager": "04",
                   "fund_portfolio": "05", "fund_costs": "05"},
    financial_streams=("fund_performance", "fund_risk"),
    readme_rows=(("01", "SID, KIM, factsheets and portfolio disclosures"),
                 ("02", "Returns, rolling returns, category rank, risk and risk-adjusted returns"),
                 ("03", "News section and every web source cited"),
                 ("04", "Fund manager, AMC and scheme events"),
                 ("05", "Portfolio, style, costs and loads"),
                 ("06", "Report (md/html/pdf), fact-check log, full claim ledger (xlsx/csv)")),
    disclaimer="Personal research, not SEBI-registered investment advice. Past returns do not guarantee future returns.",
)  # fmt: skip
BOND_LAYOUT = PackLayout(
    label="bond research report",
    folders={"01": "01_Offer_Documents", "02": "02_Issuer_Credit", "03": "03_News", "04": "04_Ratings",
             "05": "05_Pricing_and_Terms", "06": "06_Final_Report"},
    doc_folder={"OTHER": "01", "ANNUAL_REPORT": "02", "FINANCIAL_STATEMENTS": "02"},
    stream_folder={"bond_terms": "05", "bond_pricing": "05", "bond_issuer": "02", "bond_rating": "04", "bond_news": "03"},
    financial_streams=("bond_issuer",),
    readme_rows=(("01", "Offer document / information memorandum"),
                 ("02", "Issuer financials and credit analysis"),
                 ("03", "News section and every web source cited"),
                 ("04", "Rating history, drivers and sensitivities"),
                 ("05", "Terms, cash flows, yields and comparisons"),
                 ("06", "Report (md/html/pdf), fact-check log, full claim ledger (xlsx/csv)")),
    disclaimer="Personal research, not SEBI-registered investment advice. Prices and ratings are as of their dates.",
)  # fmt: skip
LAYOUTS = {"ipo_report": IPO_LAYOUT, "stock_report": STOCK_LAYOUT, "fund_report": FUND_LAYOUT,
           "bond_report": BOND_LAYOUT}  # fmt: skip


@dataclass
class PackResult:
    path: Path
    gate_ok: bool
    blocking: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    pdf_pages: int | None = None
    files: int = 0
    notes: list[str] = field(default_factory=list)


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_")[:120]


def load_claims(session, run_id: int) -> dict[int, ClaimView]:
    docs = {d.id: d.title for d in session.scalars(select(Document))}
    out: dict[int, ClaimView] = {}
    for c in session.scalars(select(Claim).where(Claim.run_id == run_id).order_by(Claim.id)):
        out[c.id] = ClaimView(
            id=c.id, status=c.status, stream=c.stream, statement=c.statement, metric=c.metric,
            value=str(c.value.normalize()) if c.value is not None else None, unit=c.unit, period=c.period,
            importance=c.importance, note=c.verifier_note, corrects=c.corrects_claim_id,
            citations=[{"doc_title": docs.get(x.document_id), "page": x.page_no, "lines": f"{x.line_start}-{x.line_end}",
                        "quote": x.quote, "quote_found": x.quote_found, "url": x.url,
                        "accessed_at": x.accessed_at.isoformat() if x.accessed_at else None} for x in c.citations],
        )  # fmt: skip
    return out


def _stream_sections(session, run_id: int) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {}
    for st in session.scalars(select(AgentStep).where(AgentStep.run_id == run_id, AgentStep.stage == "stream",
                                                      AgentStep.status == "done").order_by(AgentStep.id)):  # fmt: skip
        md = (st.output or {}).get("section_markdown")
        if md:
            label = (
                ""
                if ":" not in st.key.split("stream:")[0]
                else f"\n\n_Follow-up ({st.key.split(':')[0]})_\n\n"
            )
            sections.setdefault(st.role, []).append(label + md)
    return sections


def _latest_synthesis(session, run_id: int) -> dict[str, Any] | None:
    st = session.scalars(select(AgentStep).where(AgentStep.run_id == run_id, AgentStep.stage == "synthesis",
                                                 AgentStep.status == "done").order_by(AgentStep.finished_at.desc(),
                                                                                      AgentStep.id.desc())).first()  # fmt: skip
    return st.output if st else None


# --------------------------------------------------------------------------- tables & charts
def financial_pivot(
    claims: dict[int, ClaimView], streams: tuple[str, ...] = ("financials",)
) -> tuple[list[str], list[list[Any]]]:
    """metric x period table from usable numeric financials claims (verified first)."""
    rank = {"verified": 0, "unverified": 1, "needs_review": 2}
    cells: dict[tuple[str, str], tuple[int, str, int]] = {}
    for c in claims.values():
        if c.stream not in streams or not (c.metric and c.period and c.value) or c.status not in rank:
            continue
        key = (f"{c.metric} ({c.unit})" if c.unit else c.metric, c.period)
        cand = (rank[c.status], c.value, c.id)
        if key not in cells or cand < cells[key]:
            cells[key] = cand
    periods = sorted({p for _, p in cells})
    metrics = sorted({m for m, _ in cells})
    rows = [[m, *[(f"{cells[(m, p)][1]} [C{cells[(m, p)][2]}]" if (m, p) in cells else "") for p in periods]]
            for m in metrics]  # fmt: skip
    return ["metric", *periods], rows


def charts(claims: dict[int, ClaimView], out_dir: Path, limit: int = 4,
           streams: tuple[str, ...] = ("financials",)) -> list[Path]:  # fmt: skip
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    series: dict[str, dict[str, float]] = {}
    for c in claims.values():
        if (
            c.stream not in streams
            or c.status not in ("verified", "unverified")
            or not (c.metric and c.period and c.value)
        ):
            continue
        if not re.search(r"fy|q\d|20\d\d", c.period, re.I):
            continue
        try:
            series.setdefault(f"{c.metric} ({c.unit or ''})", {})[c.period] = float(c.value)
        except ValueError:
            continue
    picked = sorted((k for k, v in series.items() if len(v) >= 3), key=lambda k: -len(series[k]))[:limit]
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for name in picked:
        pts = sorted(series[name].items())
        fig, ax = plt.subplots(figsize=(7, 3.2), dpi=130)
        ax.bar([p for p, _ in pts], [v for _, v in pts], color="#1f5fbf")
        ax.set_title(name, fontsize=11)
        ax.grid(axis="y", alpha=0.3)
        ax.spines[["top", "right"]].set_visible(False)
        for i, (_, v) in enumerate(pts):
            ax.annotate(f"{v:,.2f}", (i, v), ha="center", va="bottom", fontsize=8)
        fig.tight_layout()
        p = out_dir / f"{_safe(name)}.png"
        fig.savefig(p)
        plt.close(fig)
        paths.append(p)
    return paths


def write_xlsx(path: Path, sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                60, max(10, *(len(str(x.value or "")) for x in col))
            )
        ws.freeze_panes = "A2"
    tmp = path.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    tmp.replace(path)


def write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> None:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    path.write_text(buf.getvalue())


def _claims_rows(claims: dict[int, ClaimView]) -> tuple[list[str], list[list[Any]]]:
    header = ["claim_id", "status", "importance", "stream", "statement", "metric", "value", "unit", "period",
              "corrects", "sources", "verifier_note"]  # fmt: skip
    rows = []
    for c in claims.values():
        src = "; ".join(x["url"] if x.get("url") else f"{x.get('doc_title')} p{x.get('page')} L{x.get('lines')}"
                        for x in c.citations)  # fmt: skip
        rows.append([c.id, c.status, c.importance, c.stream, c.statement, c.metric, c.value, c.unit, c.period,
                     c.corrects, src, (c.note or "")[:1000]])  # fmt: skip
    return header, rows


def fact_check_log(claims: dict[int, ClaimView], cited: set[int], gate_ok: bool, blocking: list[str],
                   warnings: list[str]) -> str:  # fmt: skip
    counts: dict[str, int] = {}
    for c in claims.values():
        counts[c.status] = counts.get(c.status, 0) + 1
    lines = ["# Fact-check log", "", f"Publish gate: **{'PASSED' if gate_ok else 'BLOCKED'}**", ""]
    lines += [f"- blocking: {b}" for b in blocking] + [f"- warning: {w}" for w in warnings[:40]]
    lines += ["", "## Ledger status", "", "| status | claims |", "|---|---|"]
    lines += [f"| {k} | {v} |" for k, v in sorted(counts.items())]
    lines += ["", "## Contradicted claims and corrections", "", "| claim | statement | verifier note | correction |",
              "|---|---|---|---|"]  # fmt: skip
    corr = {c.corrects: c.id for c in claims.values() if c.corrects}
    for c in claims.values():
        if c.status == "contradicted":
            fix = f"C{corr[c.id]}" if c.id in corr else ""
            lines.append(
                f"| C{c.id} | {c.statement[:140]} | {(c.note or '')[:220]} | {fix} |".replace("\n", " ")
            )
    lines += ["", "## Claims cited in the report", "", "| claim | status | statement |", "|---|---|---|"]
    lines += [f"| C{i} | {claims[i].status} | {claims[i].statement[:160]} |".replace("\n", " ")
              for i in sorted(cited) if i in claims]  # fmt: skip
    return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------- pack
def render_pack(
    run_id: int, *, out_root: Path | None = None, pdf: bool = True, copy_documents: bool = True
) -> PackResult:
    now = now_ist()
    with session_scope() as s:
        run = s.get(ResearchRun, run_id)
        if run is None:
            raise ValueError(f"unknown run {run_id}")
        co = s.get(Company, run.company_id)
        synth = _latest_synthesis(s, run_id)
        if synth is None:
            raise ValueError(f"run {run_id} has no finished synthesis to render")
        report_md = synth["report_markdown"]
        gate = check_report(s, run_id, report_md)
        claims = load_claims(s, run_id)
        sections = _stream_sections(s, run_id)
        docs = [
            (d.kind, d.title, Path(d.local_path))
            for d in s.scalars(select(Document).where(Document.company_id == co.id))
        ]
        company_name, slug = co.name, co.slug
        layout = LAYOUTS.get(run.kind, IPO_LAYOUT)

    root = (out_root or get_settings().reports_dir) / slug / f"run-{run_id}"
    folders = layout.folders
    for f in folders.values():
        (root / f).mkdir(parents=True, exist_ok=True)
    res = PackResult(path=root, gate_ok=gate.ok, blocking=gate.blocking, warnings=gate.warnings)

    if copy_documents:
        for kind, title, src in docs:
            if src.exists():
                dest = (
                    root
                    / folders[layout.doc_folder.get(kind, "05")]
                    / f"{_safe(kind)}__{_safe(title)}{src.suffix}"
                )
                if not dest.exists():
                    shutil.copyfile(src, dest)
            else:
                res.notes.append(f"source file missing for {title}")

    for stream, parts in sections.items():
        folder = root / folders[layout.stream_folder.get(stream, "05")]
        (folder / f"{stream}_section.md").write_text(f"# {stream} (run {run_id})\n\n" + "\n\n".join(parts))
    web = sorted(
        {(x["url"], x.get("accessed_at")) for c in claims.values() for x in c.citations if x.get("url")}
    )
    (root / folders["03"] / "sources.md").write_text(
        "# Web sources cited in the ledger\n\n" + "\n".join(f"- {u} (accessed {a})" for u, a in web) + "\n"
    )

    fin_header, fin_rows = financial_pivot(claims, layout.financial_streams)
    write_csv(root / folders["02"] / "financials.csv", fin_header, fin_rows)
    claim_header, claim_rows = _claims_rows(claims)
    write_xlsx(root / folders["02"] / "financials.xlsx", {"Financials": (fin_header, fin_rows)})
    write_xlsx(root / folders["06"] / "claims.xlsx", {"Claims": (claim_header, claim_rows)})
    write_csv(root / folders["06"] / "claims.csv", claim_header, claim_rows)
    chart_paths = charts(claims, root / folders["02"] / "charts", streams=layout.financial_streams)

    final = root / folders["06"]
    title = f"{company_name} — {layout.label} (run {run_id})"
    md_name = "report.md" if gate.ok else "report_NOT_PUBLISHED.md"
    (final / md_name).write_text(report_md)
    cited = set(gate.cited_claims)  # every citation spelling the gate checked
    (final / "fact_check_log.md").write_text(
        fact_check_log(claims, cited, gate.ok, gate.blocking, gate.warnings)
    )
    html = render_html(report_md, claims, title=title, generated_at=now, gate_ok=gate.ok,
                       gate_blocking=gate.blocking, gate_warnings=gate.warnings)  # fmt: skip
    if chart_paths:
        imgs = "".join(f"<figure><img alt='{p.stem}' style='max-width:100%' src='data:image/png;base64,"
                       f"{base64.b64encode(p.read_bytes()).decode()}'></figure>" for p in chart_paths)  # fmt: skip
        html = html.replace("<section class='evidence'>", f"<section><h2>Charts (from ledger claims)</h2>{imgs}</section>"
                            "<section class='evidence'>", 1)  # fmt: skip
    html_path = final / ("report.html" if gate.ok else "report_NOT_PUBLISHED.html")
    tmp = html_path.with_suffix(".html.tmp")
    tmp.write_text(html)
    tmp.replace(html_path)
    if pdf:
        try:
            res.pdf_pages = html_to_pdf(html_path, html_path.with_suffix(".pdf"), title=title)
        except PdfRenderError as e:
            res.notes.append(f"PDF not produced: {e}")

    status = (
        "PASSED — published" if gate.ok else "BLOCKED — draft only (see 06_Final_Report/fact_check_log.md)"
    )
    rows = "".join(f"| {folders[k]} | {text} |\n" for k, text in layout.readme_rows)
    (root / "README.md").write_text(
        f"# {company_name} — research pack (run {run_id})\n\n"
        f"Generated {now:%d %b %Y %H:%M} IST by FinResearch. Publish gate: **{status}**.\n\n"
        f"{layout.disclaimer}\n\n**{EXPORT_WATERMARK}.** {EXPORT_DISCLAIMER}\n\n| Folder | Contents |\n|---|---|\n{rows}"
    )
    res.files = sum(1 for p in root.rglob("*") if p.is_file())
    return res
