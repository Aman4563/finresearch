"""FBIL (Financial Benchmarks India Pvt Ltd, https://www.fbil.org.in): the G-sec par yield curve.

FBIL publishes, after each business day, "FBIL GOI Prices including Par Yield": an Excel workbook with a "Par Yield"
sheet of tenors (0.25-year steps) against the par yield, half-yearly compounded (as G-secs pay coupons) and
annualised. The site's own front end reads two endpoints:

* ``/wasdm/gsec/fetch?authenticated=false`` -> the latest published dates, newest first;
* ``/wasdm/gsec/downloadPublished?date=YYYY-MM-DD`` -> that day's workbook (served as ``.xls`` but really ``.xlsx``).

The domain is fbil.org.in. (fbil.co.in is an unrelated bullion dealer.)

When FBIL can't be reached, `FALLBACK_CURVE` is the curve published for 23-Sep-2026, labelled with its date so a
stale curve is never mistaken for today's.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import pairwise

from finresearch.adapters.http import PoliteClient

BASE = "https://www.fbil.org.in/wasdm"
LATEST_URL = f"{BASE}/gsec/fetch"
DOWNLOAD_URL = f"{BASE}/gsec/downloadPublished"
SITE = "https://www.fbil.org.in/"
HEADERS = {"Referer": SITE, "Accept": "application/json, text/plain, */*"}


class FbilError(RuntimeError):
    pass


@dataclass(frozen=True)
class ParPoint:
    tenor_years: Decimal
    semi_annual: Decimal  # fraction, compounded half-yearly (0.0666 = 6.66 %)
    annualised: Decimal | None  # fraction, FBIL's own annualised figure


@dataclass(frozen=True)
class ParCurve:
    as_of: date
    points: list[ParPoint]
    source: str = SITE
    fallback: bool = False  # True: FBIL couldn't be reached and this is the dated table below
    notes: list[str] = field(default_factory=list)

    def par_yield(self, years: float) -> Decimal:
        """Half-yearly par yield at `years` to maturity, linear between FBIL's tenors and flat beyond its ends."""
        pts = sorted(self.points, key=lambda p: p.tenor_years)
        if not pts:
            raise FbilError("empty par curve")
        y = Decimal(str(years))
        if y <= pts[0].tenor_years:
            return pts[0].semi_annual
        if y >= pts[-1].tenor_years:
            return pts[-1].semi_annual
        for a, b in pairwise(pts):
            if a.tenor_years <= y <= b.tenor_years:
                w = (y - a.tenor_years) / (b.tenor_years - a.tenor_years)
                return (a.semi_annual + (b.semi_annual - a.semi_annual) * w).quantize(Decimal("1e-8"))
        return pts[-1].semi_annual  # unreachable


def _pct(v: object) -> Decimal | None:
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v).strip()) / 100
    except InvalidOperation:
        return None


def _as_of(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v).strip(), "%d-%b-%Y").date()
    except ValueError:
        return None


def parse_par_yield_xlsx(content: bytes) -> ParCurve:
    """The "Par Yield" sheet of FBIL's G-sec workbook: a title row with the date, then 'Tenor (Year)',
    'YTM% p.a.(Semi-Annual)', 'YTM % p.a.(Annualized)' rows."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = next((ws for ws in wb.worksheets if "par" in ws.title.lower()), None)
        if sheet is None:
            raise FbilError(f"no Par Yield sheet (sheets: {[ws.title for ws in wb.worksheets]})")
        as_of, points, in_table = None, [], False
        for row in sheet.iter_rows(values_only=True):
            cells = [c for c in row if c is not None and str(c).strip() != ""]
            if not cells:
                continue
            first = str(cells[0]).strip().lower()
            if as_of is None and "par yield" in first and len(cells) > 1:
                as_of = _as_of(cells[1])
                continue
            if first.startswith("tenor"):
                in_table = True
                continue
            if in_table and len(cells) >= 2:
                try:
                    tenor = Decimal(str(cells[0]))
                except InvalidOperation:
                    continue
                semi = _pct(cells[1])
                if semi is not None:
                    points.append(ParPoint(tenor, semi, _pct(cells[2]) if len(cells) > 2 else None))
        if as_of is None or len(points) < 2:
            raise FbilError("could not read the par yield table")
        return ParCurve(as_of=as_of, points=points)
    finally:
        wb.close()


class FbilClient:
    def __init__(self, client: PoliteClient | None = None):
        self._own = client is None
        self.http = client or PoliteClient()

    async def __aenter__(self) -> FbilClient:
        return self

    async def __aexit__(self, *exc: object) -> None:
        if self._own:
            await self.http.aclose()

    async def published_dates(self) -> list[date]:
        resp = await self.http.get(
            LATEST_URL, params={"authenticated": "false"}, headers=HEADERS, cache_ttl=3600
        )
        if not resp.ok:
            raise FbilError(f"FBIL HTTP {resp.status} for {LATEST_URL}")
        days = [_iso(r.get("processRunDate")) for r in resp.json() if isinstance(r, dict)]
        return sorted((d for d in days if d), reverse=True)

    async def par_curve(self, day: date) -> ParCurve:
        resp = await self.http.get(DOWNLOAD_URL, params={"date": day.isoformat()}, headers=HEADERS,
                                   cache_ttl=30 * 86400, cache_if=lambda f: f.content[:2] == b"PK")  # fmt: skip
        if not resp.ok or resp.content[:2] != b"PK":
            raise FbilError(f"FBIL has no G-sec workbook for {day} (HTTP {resp.status})")
        curve = parse_par_yield_xlsx(resp.content)
        return ParCurve(
            as_of=curve.as_of, points=curve.points, source=f"{DOWNLOAD_URL}?date={day.isoformat()}"
        )

    async def latest_par_curve(self) -> ParCurve:
        days = await self.published_dates()
        if not days:
            raise FbilError("FBIL lists no published G-sec dates")
        return await self.par_curve(days[0])


def _iso(v: object) -> date | None:
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def _curve(rows: list[tuple[str, str, str]]) -> list[ParPoint]:
    return [ParPoint(Decimal(t), Decimal(s) / 100, Decimal(a) / 100) for t, s, a in rows]


# FBIL GSec Base/Par Yield published for 23-Sep-2026 (tenor years, % semi-annual, % annualised), read from
# https://www.fbil.org.in/wasdm/gsec/downloadPublished?date=2026-09-23 on 30-Sep-2026. Only the fallback.
_FALLBACK_ROWS = [
    ("0.25", "5.34", "5.42"), ("0.5", "5.80", "5.88"), ("0.75", "5.98", "6.07"), ("1", "6.07", "6.16"),
    ("1.5", "6.18", "6.28"), ("2", "6.33", "6.43"), ("2.5", "6.51", "6.62"), ("3", "6.66", "6.77"),
    ("3.5", "6.72", "6.83"), ("4", "6.72", "6.84"), ("4.5", "6.73", "6.85"), ("5", "6.79", "6.91"),
    ("5.5", "6.88", "7.00"), ("6", "6.93", "7.05"), ("6.5", "6.95", "7.07"), ("7", "6.96", "7.08"),
    ("7.5", "6.99", "7.11"), ("8", "7.01", "7.13"), ("8.5", "7.00", "7.12"), ("9", "7.01", "7.13"),
    ("9.5", "7.03", "7.16"), ("10", "7.07", "7.20"), ("11", "7.10", "7.23"), ("12", "7.11", "7.24"),
    ("13", "7.13", "7.25"), ("14", "7.17", "7.29"), ("15", "7.23", "7.36"), ("16", "7.29", "7.43"),
    ("17", "7.35", "7.49"), ("18", "7.41", "7.54"), ("19", "7.45", "7.59"), ("20", "7.49", "7.63"),
    ("25", "7.56", "7.71"), ("30", "7.58", "7.73"), ("35", "7.63", "7.77"), ("40", "7.64", "7.78"),
    ("50", "7.64", "7.79"),
]  # fmt: skip
FALLBACK_CURVE = ParCurve(
    as_of=date(2026, 9, 23),
    points=_curve(_FALLBACK_ROWS),
    source=f"{DOWNLOAD_URL}?date=2026-09-23",
    fallback=True,
    notes=["FBIL could not be reached: this is FBIL's curve for 23-Sep-2026, not today's"],
)
