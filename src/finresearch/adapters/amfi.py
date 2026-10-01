"""AMFI (Association of Mutual Funds in India) NAV data: the primary source for Indian mutual-fund NAVs.

* NAVAll.txt: the latest NAV of every scheme, grouped under category headings ("Open Ended Schemes(Equity Scheme -
  Mid Cap Fund)") and AMC headings ("Axis Mutual Fund").
* NAV history report: NAVs for a date range, for one AMC (`mf` code) or all AMCs.

AMFI does not publish its AMC codes, so they are discovered once by asking for a single day per code and reading
the AMC heading, then cached. A one-day history request for all AMCs gives every scheme's NAV on that day, which is
how category-peer returns are computed without downloading each peer's full history.

* TER: the total expense ratio of every scheme, regular and direct plan, per day of a month, from AMFI's TER page
  (https://www.amfiindia.com/ter-of-mf-schemes), downloaded as one Excel file per month. SEBI requires AMCs to
  publish TER daily and a direct plan's TER to be lower than its regular plan's (Master Circular for Mutual Funds).
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from finresearch.adapters.http import PoliteClient

NAV_ALL_URL = "https://www.amfiindia.com/spages/NAVAll.txt"
HISTORY_URL = "https://portal.amfiindia.com/DownloadNAVHistoryReport_Po.aspx"
MAX_AMC_CODE = 90
# The AMC-code map is discovered by probing one day per code (MAX_AMC_CODE requests). It is kept AMC_CODES_TTL_S, and
# never saved when the probe found fewer than MIN_AMC_CODES of India's ~45 AMCs: a probe day without NAVs (a weekend,
# a holiday, today before the evening publication) only finds the AMCs with liquid/overnight funds, and saving that
# map made every other AMC's history fail with "no AMFI code found" for good.
AMC_CODES_TTL_S = 30 * 86400
MIN_AMC_CODES = 20


def probe_day(today: date, holidays: set[date] | None = None) -> date:
    """The last NSE trading day before `today`: every AMC has published that day's NAVs by now. (The callers used
    "yesterday, or Friday on a Monday": a Saturday on a Sunday, a holiday after a holiday, and one used today.)"""
    from finresearch.fincalc.dates import is_business_day

    if holidays is None:
        from finresearch.adapters.nse_holidays import trading_holidays

        holidays = trading_holidays()
    d = today - timedelta(days=1)
    while not is_business_day(d, holidays):
        d -= timedelta(days=1)
    return d


TER_URL = "https://www.amfiindia.com/api/populate-te-rdata-revised"
TER_PAGE = "https://www.amfiindia.com/ter-of-mf-schemes"


class AmfiError(RuntimeError):
    pass


@dataclass(frozen=True)
class SchemeNav:
    code: str
    name: str
    plan: str | None
    option: str | None
    isin_growth: str | None
    isin_reinvest: str | None
    nav: Decimal | None
    day: date | None
    category: str | None
    amc: str | None
    structure: str | None = None  # "open", "close" or "interval" (the heading's "Open Ended" prefix)

    @property
    def is_direct_growth(self) -> bool:
        return "direct" in (self.plan or self.name).lower() and "growth" in (self.option or self.name).lower()


def _dec(v: str) -> Decimal | None:
    try:
        return Decimal(v.strip().replace(",", ""))
    except (InvalidOperation, ValueError):
        return None


def _day(v: str) -> date | None:
    try:
        return datetime.strptime(v.strip(), "%d-%b-%Y").date()
    except ValueError:
        return None


def _structure(line: str) -> str | None:
    """'Open Ended Schemes(Equity Scheme - Mid Cap Fund)' -> 'open'; 'Close Ended ...' -> 'close'; 'Interval Fund
    Schemes(...)' -> 'interval'."""
    head = line.split("(")[0].lower()
    return "open" if "open ended" in head else "close" if "close ended" in head else "interval" if "interval" in head \
        else None  # fmt: skip


def _category(line: str) -> str | None:
    """'Open Ended Schemes(Equity Scheme - Mid Cap Fund)' -> 'Equity Scheme - Mid Cap Fund'."""
    if "(" in line and line.rstrip().endswith(")") and "scheme" in line.lower().split("(")[0]:
        return line[line.index("(") + 1 : line.rindex(")")].strip()
    return None


# Column positions when a file has no header line (the layouts of the recorded files, Sep-2026)
_NAV_ALL_COLS = (
    "code",
    "g",
    "r",
    "name",
    "plan",
    "option",
    "nav",
    "day",
)  # Scheme Code;ISIN Growth;ISIN Reinvest;...
_HISTORY_COLS = (
    "code",
    "name",
    "plan",
    "option",
    "g",
    "r",
    "nav",
    "day",
)  # Scheme Code;NAV Name;Plan;Option;...


def _columns(header: list[str]) -> tuple[str, ...] | None:
    """The field of each column, from AMFI's header line ("Scheme Code;ISIN Div Payout/ ISIN Growth;ISIN Div
    Reinvestment;Scheme Name;Plan;Option;Net Asset Value;Date"; older files have no Plan/Option columns). None when a
    required column is missing."""
    out = []
    for h in (x.strip().lower() for x in header):
        if h == "scheme code":
            out.append("code")
        elif "reinvest" in h:
            out.append("r")
        elif "growth" in h or "payout" in h:
            out.append("g")
        elif h in ("scheme name", "nav name"):
            out.append("name")
        elif h in ("plan", "option", "date"):
            out.append({"date": "day"}.get(h, h))
        elif h.startswith("net asset value"):
            out.append("nav")
        else:
            out.append("")
    return tuple(out) if {"code", "name", "nav", "day"} <= set(out) else None


def _parse(text: str, history: bool) -> list[SchemeNav]:
    out: list[SchemeNav] = []
    category = amc = structure = None
    cols = _HISTORY_COLS if history else _NAV_ALL_COLS
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.lower().startswith("scheme code"):  # the header names the columns: read them, never assume
            cols = _columns(line.split(";")) or cols
            continue
        if ";" not in line:
            cat = _category(line)
            if cat is not None:
                category, structure = cat, _structure(line)
            else:
                amc = line
            continue
        f = [x.strip() for x in line.split(";")]
        if len(f) < len(cols):
            continue
        v = dict(zip(cols, f, strict=False))
        g, r = v.get("g", ""), v.get("r", "")
        out.append(SchemeNav(code=v["code"], name=v["name"], plan=v.get("plan") or None, option=v.get("option") or None,
                             isin_growth=None if g in ("", "-") else g, isin_reinvest=None if r in ("", "-") else r,
                             nav=_dec(v["nav"]), day=_day(v["day"]), category=category, amc=amc,
                             structure=structure))  # fmt: skip
    return out


def parse_nav_all(text: str) -> list[SchemeNav]:
    return _parse(text, history=False)


def parse_nav_history(text: str) -> list[SchemeNav]:
    return _parse(text, history=True)


@dataclass(frozen=True)
class SchemeTer:
    """A scheme's latest total expense ratio in a month's TER file, in percent a year (1.03 = 1.03 %)."""

    name: str
    category: str | None
    day: date
    regular: Decimal | None
    direct: Decimal | None


def ter_key(name: str) -> str:
    """Scheme names as a join key: lower case, punctuation and spacing folded ('HDFC Large Cap Fund' in the TER
    file matches the NAV file's scheme name)."""
    return re.sub(r"[^a-z0-9]+", " ", name.lower()).strip()


# AMFI's NAVAll spells the same SEBI category several ways while AMCs move to its newer headings (seen in the
# 30-Sep-2026 file: "Equity Scheme - Flexi Cap Fund" for some AMCs, "Equity Schemes - Flexi Cap Fund" for others).
# Only spellings are folded here. Headings whose *names* differ ("Debt Scheme - Short Duration Fund" vs
# "Income/Debt Oriented Schemes - Short Term Fund", "Sectoral/ Thematic" vs separate "Sectoral Fund" and
# "Thematic Fund") are kept apart: whether they are the same SEBI category is not verified here.
_GROUPS = {
    "equity scheme": "equity", "equity schemes": "equity",
    "hybrid scheme": "hybrid", "hybrid schemes": "hybrid",
    "debt scheme": "debt", "income/debt oriented schemes": "debt",
    "solution oriented scheme": "solution", "solution oriented schemes": "solution", "children's fund": "solution",
    "other scheme": "other", "index funds": "index", "exchange traded funds (etfs)": "etf",
    "fund of funds scheme (domestic)": "fof", "overseas fund of funds": "fof", "life cycle funds": "lifecycle",
}  # fmt: skip
_NAMES = {
    "elss- tax saver fund": "elss",
    "elss - tax saver fund": "elss",
    "elss tax saver fund": "elss",
    "balanced advantage fund/ dynamic asset allocation": "dynamic asset allocation or balanced advantage",
    "childrens' fund": "children's fund",
    "banking and psu debt fund": "banking and psu fund",
}


def category_key(raw: str | None) -> str | None:
    """A grouping key for AMFI category headings that folds spelling variants: "Equity Scheme - ELSS" and "Equity
    Schemes - ELSS- Tax Saver Fund" -> "equity:elss"; "Hybrid Scheme - Equity Savings" and "Hybrid Schemes - Equity
    Savings Fund" -> "hybrid:equity savings". A heading without a " - " (the legacy "Income", "Growth", "Gilt") gets
    "legacy:<name>"."""
    if not raw:
        return None
    s = re.sub(r"\s+", " ", raw.lower().replace("\u2019", "'").replace("**", "")).strip()
    if " - " not in s:
        return f"legacy:{s}"
    group, name = (x.strip() for x in s.split(" - ", 1))
    group = _GROUPS.get(group, group)
    name = _NAMES.get(name, name)
    if name.endswith(" fund"):
        name = name[: -len(" fund")]
    return f"{group}:{name}"


def parse_ter_xlsx(content: bytes) -> dict[str, SchemeTer]:
    """AMFI's monthly TER Excel file -> the latest row per scheme, keyed by `ter_key(name)`. Columns are found by
    their headings ('Scheme Name', 'Scheme Category', 'TER Date', 'Regular Plan - Total TER (%)', 'Direct Plan -
    Total TER (%)')."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        head = [str(x or "").strip().lower() for x in next(rows, ())]

        def col(*words: str) -> int:
            for i, h in enumerate(head):
                if all(w in h for w in words):
                    return i
            raise AmfiError(f"TER file has no column with {words}; headings: {head}")

        c_name, c_cat, c_day = col("scheme name"), col("scheme category"), col("ter date")
        c_reg, c_dir = col("regular", "total ter"), col("direct", "total ter")
        out: dict[str, SchemeTer] = {}
        for r in rows:
            if not r or len(r) <= max(c_name, c_reg, c_dir) or not r[c_name]:
                continue
            raw_day = r[c_day]
            day_ = (
                raw_day.date()
                if isinstance(raw_day, datetime)
                else raw_day
                if isinstance(raw_day, date)
                else None
            )
            if day_ is None:
                continue
            key = ter_key(str(r[c_name]))
            if key in out and out[key].day >= day_:
                continue
            out[key] = SchemeTer(name=str(r[c_name]).strip(), category=(str(r[c_cat]).strip() if r[c_cat] else None),
                                 day=day_, regular=_ter(r[c_reg]), direct=_ter(r[c_dir]))  # fmt: skip
        return out
    finally:
        wb.close()


def _ter(v: object) -> Decimal | None:
    if v is None or v == "":
        return None
    try:
        d = Decimal(str(v).strip())
    except InvalidOperation:
        return None
    return d if d >= 0 else None


def search_schemes(schemes: list[SchemeNav], query: str, limit: int = 20) -> list[SchemeNav]:
    """Code match first, then names containing every word; direct-growth variants first."""
    q = query.strip().lower()
    if not q:
        return []
    words = q.split()
    hits = [
        s for s in schemes if s.code == q or all(w in f"{s.name} {s.plan} {s.option}".lower() for w in words)
    ]
    return sorted(hits, key=lambda s: (s.code != q, not s.is_direct_growth, s.name))[:limit]


class AmfiClient:
    def __init__(self, client: PoliteClient | None = None, *, cache_dir: Path | None = None):
        self._own = client is None
        self.http = client or PoliteClient()
        self.cache_dir = cache_dir

    async def __aenter__(self) -> AmfiClient:
        return self

    async def __aexit__(self, *exc: object) -> None:
        if self._own:
            await self.http.aclose()

    async def _text(self, url: str, params: dict[str, str] | None = None) -> str:
        resp = await self.http.get(url, params=params)
        if not resp.ok:
            raise AmfiError(f"AMFI HTTP {resp.status} for {url}")
        return resp.content.decode("utf-8", "replace")

    async def nav_all(self) -> list[SchemeNav]:
        """Every scheme's latest NAV. A body with no NAV rows (a maintenance or block page served with HTTP 200, or a
        layout the parser does not know) raises: an empty list would read as "every fund is missing from AMFI"."""
        rows = parse_nav_all(await self._text(NAV_ALL_URL))
        if not any(r.nav is not None for r in rows):
            raise AmfiError(f"AMFI NAVAll had no NAV rows (a block page or a format change) at {NAV_ALL_URL}")
        return rows

    async def history(self, start: date, end: date, amc_code: int | None = None) -> list[SchemeNav]:
        params = {"frmdt": start.strftime("%d-%b-%Y"), "todt": end.strftime("%d-%b-%Y")}
        if amc_code is not None:
            params["mf"] = str(amc_code)
        return parse_nav_history(await self._text(HISTORY_URL, params))

    async def ter(self, month: date) -> dict[str, SchemeTer]:
        """Every scheme's latest TER in `month` (AMFI's Excel download; cached on disk for a day)."""
        params = {
            "MF_ID": "All",
            "Month": month.strftime("%m-%Y"),
            "strCat": "-1",
            "strType": "-1",
            "excel": "true",
        }
        resp = await self.http.get(TER_URL, params=params, cache_ttl=86400,
                                   cache_if=lambda f: f.content[:2] == b"PK")  # fmt: skip
        if not resp.ok or resp.content[:2] != b"PK":
            raise AmfiError(f"AMFI TER download failed (HTTP {resp.status}) for {month:%m-%Y}")
        return parse_ter_xlsx(resp.content)

    async def navs_on(self, day: date, lookback_days: int = 7) -> dict[str, SchemeNav]:
        """Every scheme's latest NAV on or before `day` (a holiday falls back to the previous NAV date)."""
        rows = await self.history(day - timedelta(days=lookback_days), day)
        latest: dict[str, SchemeNav] = {}
        for r in rows:
            if (
                r.day
                and r.day <= day
                and r.nav is not None
                and (r.code not in latest or r.day > latest[r.code].day)
            ):
                latest[r.code] = r
        return latest

    def _codes_cache(self) -> Path | None:
        return self.cache_dir / "amfi_amc_codes.json" if self.cache_dir else None

    def _codes_age_s(self) -> float | None:
        import time

        cache = self._codes_cache()
        return time.time() - cache.stat().st_mtime if cache and cache.exists() else None

    async def amc_codes(self, probe_day: date, *, refresh: bool = False) -> dict[str, int]:
        """AMC name -> AMFI `mf` code, discovered by asking for one day per code (cached in cache_dir for
        AMC_CODES_TTL_S, and only when the probe found at least MIN_AMC_CODES AMCs)."""
        cache = self._codes_cache()
        age = self._codes_age_s()
        if cache and age is not None and age < AMC_CODES_TTL_S and not refresh:
            return json.loads(cache.read_text())
        codes: dict[str, int] = {}
        for code in range(1, MAX_AMC_CODE + 1):
            rows = await self.history(probe_day, probe_day, amc_code=code)
            amc = next((r.amc for r in rows if r.amc), None)
            if amc:
                codes[amc] = code
        if cache and len(codes) >= MIN_AMC_CODES:
            cache.parent.mkdir(parents=True, exist_ok=True)
            cache.write_text(json.dumps(codes, indent=1, sort_keys=True))
        return codes

    async def scheme_history(
        self, scheme: SchemeNav, start: date, end: date, probe_day: date
    ) -> list[SchemeNav]:
        codes = await self.amc_codes(probe_day)
        age = self._codes_age_s()
        if scheme.amc not in codes and age is not None and age > 86400:
            codes = await self.amc_codes(
                probe_day, refresh=True
            )  # a new or renamed AMC: re-probe, at most daily
        if scheme.amc not in codes:
            raise AmfiError(f"no AMFI code found for {scheme.amc!r}")
        rows = await self.history(start, end, amc_code=codes[scheme.amc])
        return sorted((r for r in rows if r.code == scheme.code and r.nav is not None and r.day is not None),
                      key=lambda r: r.day)  # fmt: skip
