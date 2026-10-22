"""Read an Excel workbook from bytes as plain rows, whichever format the server sent (#366).

Official sources still serve the old OLE2 `.xls` format: FBIL's G-sec workbooks up to 2022 (later ones are `.xlsx`
under an `.xls` name) and NSE's fii_stats (#371). The format is told by the body's magic bytes, never by the name or
the date: `PK` (a zip) is `.xlsx` and goes to openpyxl; `D0 CF 11 E0` (OLE2) is `.xls` and goes to xlrd.

The .xls reader is xlrd 2.0.2 (decision in #366): BSD-3 (compatible with this MIT repo), pure Python with no
dependencies (no native code parsing downloaded files), "Production/Stable", the reader pandas uses for .xls, and
since 2.0 it reads only .xls, which is exactly what is needed. 2.0.2 was released 14-Jun-2025. python-calamine (MIT,
Rust) was the alternative considered: actively developed but "Beta", a compiled wheel, and its newest release was
days old when this was decided.

Cells come back as Python values: text as str, numbers as float (int for an .xlsx integer), dates as datetime, empty
as None; trailing empty cells are dropped from each row, so an empty row is [].
"""

from __future__ import annotations

import io
from typing import Any

OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK"


class NotAWorkbook(ValueError):
    """The body is neither an .xls nor an .xlsx workbook (a block page, an HTML error, a truncated download)."""


def kind(content: bytes) -> str | None:
    if content[:8] == OLE2_MAGIC:
        return "xls"
    if content[:2] == ZIP_MAGIC:
        return "xlsx"
    return None


def _trim(row: list[Any]) -> list[Any]:
    while row and (row[-1] is None or (isinstance(row[-1], str) and not row[-1].strip())):
        row.pop()
    return row


def _read_xls(content: bytes) -> dict[str, list[list[Any]]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        out: dict[str, list[list[Any]]] = {}
        for name in book.sheet_names():
            sh = book.sheet_by_name(name)
            rows = []
            for r in range(sh.nrows):
                vals: list[Any] = []
                for c in sh.row(r):
                    if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        vals.append(None)
                    elif c.ctype == xlrd.XL_CELL_DATE:
                        vals.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                    elif c.ctype == xlrd.XL_CELL_ERROR:
                        vals.append(None)  # #N/A, #DIV/0!: unknown, never 0
                    else:
                        vals.append(c.value)
                rows.append(_trim(vals))
            out[name] = rows
            book.unload_sheet(name)
        return out
    finally:
        book.release_resources()


def _read_xlsx(content: bytes) -> dict[str, list[list[Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return {ws.title: [_trim(list(r)) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    finally:
        wb.close()


def read_sheets(content: bytes) -> dict[str, list[list[Any]]]:
    """{sheet name: rows} in workbook order. Raises NotAWorkbook for anything else, and for a workbook the reader
    cannot open (the reader's own error is kept as the cause)."""
    k = kind(content)
    if k is None:
        raise NotAWorkbook(f"not an Excel workbook (starts {content[:8]!r})")
    try:
        return _read_xls(content) if k == "xls" else _read_xlsx(content)
    except Exception as e:
        raise NotAWorkbook(f"unreadable {k} workbook: {type(e).__name__}: {e}"[:300]) from e
