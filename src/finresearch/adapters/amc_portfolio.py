"""Monthly scheme portfolios published by mutual-fund houses (AMCs): parse, discover, download and cache.

Why these files exist. SEBI Master Circular for Mutual Funds (as on 20-Mar-2026), para 6.1.1: "Mutual Funds/ AMCs
shall disclose portfolio (along with ISIN) as on the last day of the month for all their schemes on their respective
website and on the website of AMFI within 10 calendar days from the close of each month in a user-friendly and
downloadable spreadsheet format", in Format No. 4C (debt schemes also fortnightly). [V: circular text read
30-Sep-2026, https://www.sebi.gov.in master circular PDF, page 85.]

One parser for every AMC. Format 4C gives every house the same columns (name of the instrument, ISIN, industry or
rating, quantity, market value in Rs lakh, % to net assets); they differ in column order, header wording, extra code
columns, how many schemes share one workbook (one sheet per scheme), and the wrapper (xlsx, an xlsx named .xls, a
zip of xlsx). So the parser finds the header row by its words, maps each column by header text (never by position),
and walks the rows until GRAND TOTAL. Per-AMC differences are header aliases, not code paths.

Units are fixed here, once: `weight` is always PERCENT OF NET ASSETS (7.63 means 7.63 %). The files store fractions
(0.0763) or percents; the scale is decided per sheet from the GRAND TOTAL row (1 -> fractions, 100 -> percent).

What the files can and cannot tell (limits, stated in the UI):
- month-end snapshots only (window dressing is possible), published up to 10 days later;
- "Arbitrage" positions (long stock + short future) are hedged: they are kept out of equity exposure and overlap.
  Not every AMC labels them separately; where it does not, hedged longs count as equity [limit];
- derivatives, TREPS, cash and net receivables have no ISIN; they are the remainder "cash & others".

Where the files are (checked 30-Sep-2026, see AMC_SOURCES): PPFAS, Nippon India and DSP publish static links on a
plain HTML page, so the app can discover and download them. Axis publishes static, public file URLs but lists them
through an API that needs a site token, so Axis files come in by pasted URL or upload. Any other AMC: upload.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any
from urllib.parse import urljoin, urlsplit

ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")
MAX_FILE_BYTES = 25 * 1024 * 1024
# an xlsx (a zip of XML) or a zip of them, inflated: a portfolio workbook is a few MB of XML; a zip bomb is gigabytes
MAX_INFLATED_BYTES = 400 * 1024 * 1024


def _check_inflated(z: zipfile.ZipFile, label: str) -> None:
    """Refuse an archive whose members would inflate beyond MAX_INFLATED_BYTES (declared sizes; zipfile stops at the
    declared size, so a lying header fails its CRC check instead of inflating further)."""
    if sum(i.file_size for i in z.infolist()) > MAX_INFLATED_BYTES:
        raise AmcPortfolioError(f"{label}: would inflate beyond {MAX_INFLATED_BYTES // (1024 * 1024)} MB")


# kinds of holding; only EQUITY_KINDS enter overlap, look-through stock exposure and active share
EQUITY_KINDS = frozenset({"equity", "foreign_equity"})
KIND_LABEL = {
    "equity": "Indian equity",
    "foreign_equity": "Foreign equity",
    "arbitrage": "Arbitrage (hedged equity)",
    "reit_invit": "REITs & InvITs",
    "debt": "Bonds & money market",
    "govt": "Government securities",
    "mf_units": "Mutual fund units",
    "other": "Other instruments",
}


class AmcPortfolioError(ValueError):
    """The file is not a monthly portfolio this parser can read (the message says what to do)."""


# ----------------------------------------------------------------------------------------------- data
@dataclass
class Holding:
    isin: str
    name: str
    industry: str | None
    quantity: Decimal | None
    value_lakh: Decimal | None
    weight: Decimal  # percent of net assets
    kind: str
    section: str = ""

    def to_json(self) -> dict[str, Any]:
        d = asdict(self)
        for k in ("quantity", "value_lakh", "weight"):
            d[k] = None if d[k] is None else str(d[k])
        return d

    @staticmethod
    def from_json(d: dict[str, Any]) -> Holding:
        return Holding(isin=d["isin"], name=d["name"], industry=d.get("industry"),
                       quantity=None if d.get("quantity") is None else Decimal(d["quantity"]),
                       value_lakh=None if d.get("value_lakh") is None else Decimal(d["value_lakh"]),
                       weight=Decimal(d["weight"]), kind=d["kind"], section=d.get("section") or "")  # fmt: skip


@dataclass
class SchemePortfolio:
    sheet: str
    scheme_name: str
    as_of: date | None
    holdings: list[Holding]
    benchmark: str | None = None
    grand_total_lakh: Decimal | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def key(self) -> str:
        return scheme_key(self.scheme_name)

    @property
    def month(self) -> str | None:
        return self.as_of.strftime("%Y-%m") if self.as_of else None

    def weight_of(self, *kinds: str) -> Decimal:
        return sum((h.weight for h in self.holdings if h.kind in kinds), Decimal(0))

    def to_json(self) -> dict[str, Any]:
        return {"sheet": self.sheet, "scheme_name": self.scheme_name, "key": self.key,
                "as_of": self.as_of.isoformat() if self.as_of else None, "benchmark": self.benchmark,
                "grand_total_lakh": None if self.grand_total_lakh is None else str(self.grand_total_lakh),
                "warnings": self.warnings, "holdings": [h.to_json() for h in self.holdings]}  # fmt: skip

    @staticmethod
    def from_json(d: dict[str, Any]) -> SchemePortfolio:
        return SchemePortfolio(sheet=d["sheet"], scheme_name=d["scheme_name"],
                               as_of=date.fromisoformat(d["as_of"]) if d.get("as_of") else None,
                               holdings=[Holding.from_json(h) for h in d.get("holdings", [])],
                               benchmark=d.get("benchmark"),
                               grand_total_lakh=Decimal(d["grand_total_lakh"]) if d.get("grand_total_lakh") else None,
                               warnings=list(d.get("warnings") or []))  # fmt: skip


# ----------------------------------------------------------------------------------------------- names
_PLAN_WORDS = re.compile(r"\b(direct|regular|plan|growth|option|idcw|dividend|payout|reinvestment|bonus|"
                         r"institutional)\b")  # fmt: skip
_JOINED = [("mid cap", "midcap"), ("small cap", "smallcap"), ("large cap", "largecap"), ("flexi cap", "flexicap"),
           ("multi cap", "multicap"), ("large & mid", "large and mid")]  # fmt: skip


def clean_scheme_name(raw: str) -> str:
    """'Parag Parikh Flexi Cap Fund (An open-ended dynamic equity scheme ...)' -> 'Parag Parikh Flexi Cap Fund'."""
    s = re.sub(r"\s+", " ", str(raw or "")).strip()
    s = re.split(
        r"\s*\((?:an?\s+open|a\s+close|an?\s+interval|open[\s-]ended|close[\s-]ended)", s, flags=re.I
    )[0]
    return s.strip(" -–")


def scheme_key(name: str) -> str:
    """A plan- and case-insensitive key for one scheme: every plan/option of a scheme shares its portfolio."""
    s = clean_scheme_name(name).lower().replace("&", " & ")
    s = re.split(r"\s+-\s+(?=(?:direct|regular|growth|idcw)\b)", s)[0]
    for a, b in _JOINED:
        s = s.replace(a, b)
    s = s.replace("&", " and ")
    words = re.sub(r"[^a-z0-9]+", " ", s).split()
    while words and _PLAN_WORDS.fullmatch(
        words[-1]
    ):  # only trailing plan words: "Growth Mid Cap Fund" keeps its own
        words.pop()
    return " ".join(words)


# ----------------------------------------------------------------------------------------------- parsing helpers
def _num(v: Any) -> Decimal | None:
    """A cell as a number: floats via str (no binary noise); '$0.00%', '1,234.5', '*' (DSP: < 0.01 %) handled."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).strip().replace(",", "").replace("$", "").replace("₹", "")
    if s in ("*", "-", "--"):
        return Decimal(0) if s == "*" else None
    pct = s.endswith("%")
    s = s.rstrip("%").strip()
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        d = Decimal(s)
    except (InvalidOperation, ValueError):
        return None
    return d / 100 if pct else d  # '0.05%' in a fraction-scaled column is a fraction like the others


def _is_pct_text(v: Any) -> bool:
    """A text cell written with a percent sign ('7.63%'): `_num` returns it as a fraction (0.0763)."""
    return isinstance(v, str) and v.strip().endswith("%")


def _industry(v: Any) -> str | None:
    """'Computer Software: Prepackaged Software ##' -> without the footnote markers (#, *, ^, ~); quant writes 'N.A.'
    in its RATING/INDUSTRY columns for rows without one: that is no label, not a sector called "N.A."."""
    t = re.sub(r"[\s#*^~$@]+$", "", _text(v)).strip()
    return None if not t or re.fullmatch(r"n\.?\s*a\.?|-+|nil", t, re.I) else t


def _text(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov",
                                       "dec"], start=1)}  # fmt: skip
_DATE_RES = [
    re.compile(
        r"as\s+on\s+([A-Za-z]+)\s*(\d{1,2})(?:st|nd|rd|th)?\s*,?\s*(\d{4})", re.I
    ),  # August 31, 2026 / 31,2026
    re.compile(
        r"as\s+(?:on|at)\s+(\d{1,2})(?:st|nd|rd|th)?[\s-]+([A-Za-z]+)[\s,-]+(\d{4})", re.I
    ),  # 31 Aug 2026
    re.compile(
        r"as\s+(?:on|at)\s+(\d{1,2})[/.-](\d{1,2})[/.-](\d{4}|\d{2})\b", re.I
    ),  # 31/08/2026, Tata 31-08-26
]


def parse_as_of(text: str) -> date | None:
    for i, rx in enumerate(_DATE_RES):
        m = rx.search(text or "")
        if not m:
            continue
        try:
            if i == 0:
                return date(int(m.group(3)), _MONTHS[m.group(1)[:3].lower()], int(m.group(2)))
            if i == 1:
                return date(int(m.group(3)), _MONTHS[m.group(2)[:3].lower()], int(m.group(1)))
            y = int(m.group(3))
            return date(y + 2000 if y < 100 else y, int(m.group(2)), int(m.group(1)))
        except (KeyError, ValueError):
            continue
    return None


# header aliases: SEBI Format 4C wording as each AMC prints it (PPFAS, Axis, Nippon India, DSP files of Aug-2026;
# quant and Tata files of Aug-2026, read 06-Oct-2026: 'ISIN CODE', 'MKT VAL(Rs. Lacs)', separate RATING and INDUSTRY)
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("name of the instrument", "name of instrument", "name", "instrument name", "security name", "name of security",
             "company name", "issuer"),
    "isin": ("isin",),
    "industry": ("industry / rating", "industry/rating", "rating/industry", "rating / industry", "industry",
                 "sector", "rating"),
    "quantity": ("quantity", "no. of shares", "units"),
    "value": ("market/fair value", "market value", "fair value", "market/ fair value", "mkt val"),
    "weight": ("% to net assets", "% to nav", "% of net assets", "% to aum", "% of nav", "% to net asset",
               "percentage to nav", "% to net assets", "% net assets"),
}  # fmt: skip


def _header_map(row: tuple) -> dict[str, int] | None:
    """Column index by field for a header row, or None when the row is not the holdings header."""
    cols: dict[str, int] = {}
    for i, c in enumerate(row):
        t = _text(c).lower()
        if not t:
            continue
        # quant prints RATING before INDUSTRY: an industry/sector column wins over a rating-only one
        if "industry" in cols and t.startswith(("industry", "sector")) and _text(row[cols["industry"]]).lower(
        ).startswith("rating"):  # fmt: skip
            cols["industry"] = i
            continue
        for fld, aliases in HEADER_ALIASES.items():
            if fld in cols:
                continue
            if any(t.startswith(a) or t == a for a in aliases) or (
                fld == "weight" and "%" in t and "net" in t
            ):
                cols[fld] = i
                break
    return cols if {"isin", "weight", "name"} <= set(cols) else None


_STOP = re.compile(r"^(grand total|net assets?$|total net assets)", re.I)
_TOTAL = re.compile(r"^(sub\s*-?\s*total|total)\b|\btotal$", re.I)  # Tata: 'EQUITY & EQUITY RELATED TOTAL'
_FUNDLIKE = re.compile(r"\b(fund|etf|fof|plan|scheme|bees)\b", re.I)
_BLURB = re.compile(r"^[^A-Za-z0-9]|[•:]|^(an?|this|investors?|investment|long term)\b|mutual fund$", re.I)
# quant: 'Total Exposure due to futures (non hedging positions) as a %age of net assets' (25.3 in Aug-2026 Flexi Cap)
_NON_HEDGE = re.compile(r"exposure.*futures.*non[\s-]*hedg|non[\s-]*hedg.*exposure", re.I)


def _section(label: str, current: str) -> str:
    """The section a label row opens (sticky until the next label that names another section)."""
    t = label.lower()
    if re.search(r"\barbitrage\b", t):
        return "arbitrage"
    if (re.search(r"foreign|overseas|international|adr|gdr", t) and "equity" in t) or "foreign securit" in t:
        return "foreign_equity"
    if re.search(r"\breits?\b|\binvits?\b|real estate investment trust|infrastructure investment trust", t):
        return "reit_invit"
    if re.search(r"government securities|treasury bill|t-bill|state development|\bsdl\b|\bg-?sec", t):
        return "govt"
    if re.search(r"money market|certificate of deposit|commercial paper|debt instrument|bonds?\b|ncd|debenture|"
                 r"securitised|pass through|triparty|treps|\brepo\b", t):  # fmt: skip
        return "debt"
    if re.search(r"mutual fund|units of|exchange traded fund|\betf\b|fund units", t):
        return "mf_units"
    if re.search(r"equity|shares", t):
        return "equity"
    return current


def _kind(isin: str, section: str) -> str:
    """Section first (the file says what the row is); ISIN structure only when the section says nothing."""
    domestic = isin.startswith("IN")
    if section == "arbitrage":
        return "arbitrage"
    if section in ("equity", "foreign_equity"):
        if isin.startswith(
            "INF"
        ):  # Tata lists ETF units under "Equity & equity related": fund units, not a stock
            return "mf_units"
        return "equity" if domestic else "foreign_equity"
    if section:
        return section
    if not domestic:
        return "foreign_equity"
    if isin.startswith("INF"):
        return "mf_units"
    if isin.startswith("INE") and isin[7:9] == "01":  # INE + issuer(4) + security type '01' = equity shares
        return "equity"
    if re.match(r"^IN[0-9]", isin):
        return "govt"
    return "debt"


_BENCH_RES = [
    re.compile(r"benchmark\s*(?:name|index)?\s*[:\-–]\s*(.+)$", re.I),
    re.compile(r"benchmark\s+riskometer\s*[:\-–]\s*(.+)$", re.I),
    re.compile(r"benchmark(?:'s)?\s+risk-?o-?meter\s*[:\-–]\s*(.+)$", re.I),
]


def _benchmark(rows: list[tuple], start: int) -> str | None:
    """The scheme's benchmark as the file names it ('Benchmark Name - BSE MIDCAP 150 TRI', 'Benchmark Riskometer:
    Nifty 500 TRI', PPFAS: '(Nifty 500 TRI)' under "AMFI Tier 1 Benchmark's Riskometer")."""
    for r_i in range(start, len(rows)):
        for c_i, c in enumerate(rows[r_i]):
            t = _text(c)
            if "enchmark" not in t and "ENCHMARK" not in t:
                continue
            for rx in _BENCH_RES:
                m = rx.search(t)
                if m and len(m.group(1).strip()) > 3:
                    return m.group(1).strip()
            if re.search(r"benchmark'?s?\s+risk", t, re.I) and r_i + 1 < len(rows):
                below = rows[r_i + 1]
                for cand in (below[c_i] if c_i < len(below) else None, *below):
                    s = _text(cand)
                    if s.startswith("(") and s.endswith(")") and len(s) > 4:
                        return s[1:-1].strip()
    return None


def parse_sheet(rows: list[tuple], sheet: str) -> SchemePortfolio | None:
    """One scheme's portfolio from one sheet's rows, or None when the sheet has no holdings table."""
    header_at = cols = None
    for i, row in enumerate(rows[:40]):
        m = _header_map(row)
        if m:
            header_at, cols = i, m
            break
    if header_at is None or cols is None:
        return None
    top_texts = [_text(c) for r in rows[:header_at] for c in r if _text(c)]
    as_of = next((d for t in top_texts if (d := parse_as_of(t))), None)
    names = [t for t in top_texts if not re.search(r"portfolio|statement|as on|as at|^index$", t, re.I)
             and not re.fullmatch(r"[A-Z0-9_]{2,12}", t)]  # fmt: skip
    # the longest title line, among lines that read like a scheme name: quant and Tata put the scheme's description
    # and SEBI's suitability blurb ("*Investors should consult ...", "• Long term ...") above the table
    fundlike = [
        n for t in names if (n := clean_scheme_name(t)) and _FUNDLIKE.search(n) and not _BLURB.search(n)
    ]
    scheme_name = (
        max(fundlike, key=len) if fundlike else clean_scheme_name(max(names, key=len)) if names else sheet
    )

    def cell(row: tuple, fld: str) -> Any:
        i = cols.get(fld)
        return row[i] if i is not None and i < len(row) else None

    holdings: list[Holding] = []
    section, grand = "", None
    grand_weight: Decimal | None = None
    pct_text: set[int] = set()  # holdings whose weight cell was '7.63%' text
    end = len(rows)
    for r_i in range(header_at + 1, len(rows)):
        row = rows[r_i]
        texts = [_text(c) for c in row if _text(c)]
        if not texts:
            continue
        first = texts[0] if not ISIN_RE.match(texts[0]) else ""
        label = next((t for t in texts if not re.fullmatch(r"[A-Z0-9_]{2,14}|\d+", t)), first)
        if _STOP.match(label):
            grand, grand_weight = _num(cell(row, "value")), _num(cell(row, "weight"))
            if grand_weight is not None and _is_pct_text(cell(row, "weight")):
                grand_weight *= 100  # '100.00%' text is a percent total, not the fraction 1.00
            if grand_weight is None:  # the weight may sit in another column on the total row
                nums = [n for c in row if (n := _num(c)) is not None]
                grand_weight = nums[-1] if nums else None
            end = r_i
            break
        if r_i > header_at + 1 and (again := _header_map(row)):
            if (again["isin"], again["weight"]) == (cols["isin"], cols["weight"]):
                cols = again  # Tata repeats the header above its debt block (RATINGS for INDUSTRY): carry on
                continue
            end = r_i  # a second table (derivatives) starts: holdings are over
            break
        isin = _text(cell(row, "isin")).upper()
        if not ISIN_RE.match(isin):
            if not _TOTAL.search(label):
                section = _section(label, section)
            continue
        w = _num(cell(row, "weight"))
        if w is None:
            continue
        if _is_pct_text(cell(row, "weight")):
            pct_text.add(len(holdings))  # already a fraction of 1 whatever the column's scale
        holdings.append(Holding(isin=isin, name=_text(cell(row, "name")) or isin,
                                industry=_industry(cell(row, "industry")), quantity=_num(cell(row, "quantity")),
                                value_lakh=_num(cell(row, "value")), weight=w, kind=_kind(isin, section),
                                section=section))  # fmt: skip
    if not holdings:
        return None
    warnings: list[str] = []
    # scale: GRAND TOTAL says 1 (fractions) or 100 (percent); without it, the size of the weights decides
    # infer from the plain-number cells only: '%' text cells say nothing about the column's own scale
    raw_sum = sum((h.weight for i, h in enumerate(holdings) if i not in pct_text), Decimal(0))
    if grand_weight is not None and grand_weight > 0:
        fractions = grand_weight <= Decimal("1.5")
    else:
        fractions = raw_sum <= Decimal("1.5")
        warnings.append("no GRAND TOTAL row: the weight scale was inferred from the weights")
    # '%' text cells were read as fractions by `_num`: they become percent with the rest in a fraction-scaled
    # column, and on their own in a percent-scaled one (else a mixed sheet reads them 100x too small)
    for i, h in enumerate(holdings):
        if fractions or i in pct_text:
            h.weight *= 100
    total = sum((h.weight for h in holdings), Decimal(0))
    if total > Decimal("110") or total < Decimal("1"):
        warnings.append(f"ISIN holdings add up to {total:.2f} % of net assets: check the file")
    for row in rows[
        end:
    ]:  # unhedged stock futures have no ISIN line: their exposure is outside look-through equity
        texts = [_text(c) for c in row if _text(c)]
        if texts and _NON_HEDGE.search(texts[0]):
            got = next((n for c in row[1:] if not isinstance(c, str) and (n := _num(c)) is not None), None)
            if got:
                warnings.append(f"{got} % of net assets in non-hedging derivative positions (stock futures): this "
                                "exposure has no ISIN line, so it is not in the look-through equity")  # fmt: skip
            break
    return SchemePortfolio(sheet=sheet, scheme_name=scheme_name, as_of=as_of, holdings=holdings,
                           benchmark=_benchmark(rows, end), grand_total_lakh=grand, warnings=warnings)  # fmt: skip


def _xlsx_portfolios(data: bytes, label: str) -> list[SchemePortfolio]:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            _check_inflated(z, label)
    except zipfile.BadZipFile as e:
        raise AmcPortfolioError(f"{label}: not a readable Excel workbook (BadZipFile)") from e
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise AmcPortfolioError(f"{label}: not a readable Excel workbook ({type(e).__name__})") from e
    out = []
    try:
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            p = parse_sheet(rows, ws.title)
            if p is not None:
                out.append(p)
    finally:
        wb.close()
    return out


def parse_file(data: bytes, filename: str = "file") -> list[SchemePortfolio]:
    """Every scheme portfolio in an uploaded or downloaded file. Detects the format from the bytes, not the name:
    xlsx (also Nippon's xlsx saved as .xls), a zip of xlsx files (DSP), or legacy .xls (BIFF), which needs a re-save
    as .xlsx (no extra dependency for a format the AMCs are leaving)."""
    if len(data) > MAX_FILE_BYTES:
        raise AmcPortfolioError(f"{filename}: larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise AmcPortfolioError(f"{filename}: an old-format .xls workbook. Open it in Excel, Numbers or LibreOffice, "
                                "save it as .xlsx and upload that.")  # fmt: skip
    if data[:2] != b"PK":
        raise AmcPortfolioError(f"{filename}: not an Excel (.xlsx) file or a zip of them")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        _check_inflated(z, filename)
        names = z.namelist()
        if "[Content_Types].xml" in names:  # an xlsx workbook itself
            out = _xlsx_portfolios(data, filename)
        else:  # a zip of workbooks
            out = []
            for n in names:
                if n.lower().endswith((".xlsx", ".xlsm", ".xls")) and not n.startswith("__MACOSX"):
                    body = z.read(n)
                    if body[:2] == b"PK":
                        out += _xlsx_portfolios(body, n)
    if not out:
        raise AmcPortfolioError(f"{filename}: no portfolio table found (a header row with ISIN, name and "
                                "'% to net assets' columns)")  # fmt: skip
    return out


# ----------------------------------------------------------------------------------------------- AMFI cap list
AMFI_CAP_LIST_URLS = [  # newest first; the Jun-2026 file 404'd on 30-Sep-2026, Dec-2025 downloaded
    "https://www.amfiindia.com/Themes/Theme1/downloads/AverageMarketCapitalization30Jun2026.xlsx",
    "https://www.amfiindia.com/Themes/Theme1/downloads/AverageMarketCapitalization31Dec2025.xlsx",
]


def parse_amfi_cap_list(data: bytes) -> dict[str, Any]:
    """AMFI's six-monthly average market-cap list: {"as_of", "title", "by_isin": {isin: [rank, category]}}.
    Category is AMFI's own column ('Large Cap' / 'Mid Cap' / 'Small Cap' per SEBI's 1-100 / 101-250 / 251+)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        rows = [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    title = _text(rows[0][0]) if rows and rows[0] else ""
    head_i = next(i for i, r in enumerate(rows[:10]) if any(_text(c).lower() == "isin" for c in r))
    head = [_text(c).lower() for c in rows[head_i]]
    i_isin = head.index("isin")
    i_cat = next(i for i, h in enumerate(head) if h.startswith("categorization") or h.startswith("category"))
    i_rank = next((i for i, h in enumerate(head) if h.startswith("sr")), 0)
    out: dict[str, list] = {}
    for r in rows[head_i + 1 :]:
        isin = _text(r[i_isin]).upper() if i_isin < len(r) else ""
        cat = _text(r[i_cat]) if i_cat < len(r) else ""
        if ISIN_RE.match(isin) and cat:
            rank = _num(r[i_rank])
            out[isin] = [int(rank) if rank is not None else None, cat]
    m = re.search(r"ended\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", title)
    as_of = None
    if m:
        try:
            as_of = date(int(m.group(3)), _MONTHS[m.group(2)[:3].lower()], int(m.group(1))).isoformat()
        except (KeyError, ValueError):
            as_of = None
    return {"as_of": as_of, "title": title, "by_isin": out}


# ----------------------------------------------------------------------------------------------- where the files are
@dataclass(frozen=True)
class AmcSource:
    amc: str  # as in AMFI NAVAll's AMC heading
    page: str
    mode: str  # "auto": the app finds the monthly file links on `page`; "url": paste a file URL; upload always works
    note: str
    hosts: tuple[str, ...]


AMC_SOURCES: dict[str, AmcSource] = {
    "ppfas": AmcSource("PPFAS Mutual Fund", "https://amc.ppfas.com/downloads/portfolio-disclosure/", "auto",
                       "One xlsx per scheme per month (static links; files before Mar-2026 are old .xls).",
                       ("amc.ppfas.com",)),
    "nippon": AmcSource("Nippon India Mutual Fund",
                        "https://mf.nipponindiaim.com/investor-service/downloads/factsheet-portfolio-and-other-disclosures",
                        "auto", "One workbook for all schemes per month (NIMF-MONTHLY-PORTFOLIO-…, an xlsx named .xls).",
                        ("mf.nipponindiaim.com",)),
    "dsp": AmcSource("DSP Mutual Fund", "https://www.dspim.com/mandatory-disclosures/portfolio-disclosures", "auto",
                     "One zip per month with an equity and a debt workbook.", ("www.dspim.com",)),
    "axis": AmcSource("Axis Mutual Fund", "https://www.axismf.com/statutory-disclosures", "url",
                      "Statutory disclosures → Portfolios → Monthly Scheme Portfolios. The file links are public but "
                      "the list needs the site's own token, so copy the file's link (or download it) and paste/upload "
                      "it here.", ("www.axismf.com",)),
}  # fmt: skip
ALLOWED_HOSTS = frozenset(h for s in AMC_SOURCES.values() for h in s.hosts) | {
    "www.amfiindia.com",
    "portal.amfiindia.com",
}


def source_for_amc(amc: str | None, scheme_name: str = "") -> str | None:
    t = f"{amc or ''} {scheme_name}".lower()
    for key, words in (("ppfas", ("ppfas", "parag parikh")), ("nippon", ("nippon",)), ("dsp", ("dsp ",)),
                       ("axis", ("axis ",))):  # fmt: skip
        if any(w in t + " " for w in words):
            return key
    return None


def check_url(url: str) -> str:
    """Only https links on the AMC/AMFI hosts above: the fetch endpoint must not become a way to reach anything
    else from this machine."""
    parts = urlsplit(url.strip())
    if parts.scheme != "https" or (parts.hostname or "").lower() not in ALLOWED_HOSTS:
        raise AmcPortfolioError("only https links on these hosts can be fetched: " + ", ".join(sorted(ALLOWED_HOSTS))
                                + ". For any other fund house, download the file and upload it.")  # fmt: skip
    return url.strip()


@dataclass(frozen=True)
class FileLink:
    url: str
    month: str  # YYYY-MM of the portfolio date
    label: str  # scheme prefix (PPFAS) or "all"


_MONTH_WORD = r"(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sep(?:t(?:ember)?)?|" \
              r"oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"  # fmt: skip


def _ym(month_word: str, year: str) -> str:
    y = int(year) + (2000 if len(year) == 2 else 0)
    return f"{y:04d}-{_MONTHS[month_word[:3].lower()]:02d}"


def discover_links(source: str, html: str, page_url: str) -> list[FileLink]:
    """Monthly portfolio file links on an AMC's disclosure page, newest month first."""
    links: list[FileLink] = []
    if source == "ppfas":
        for m in re.finditer(r'href="(/downloads/portfolio-disclosure/\d{4}/([A-Z]+)_PPFAS_Monthly_Portfolio_Report_'
                             r'([A-Za-z]+)_\d{1,2}_(\d{4})\.xlsx)[^"]*"', html):  # fmt: skip
            links.append(FileLink(urljoin(page_url, m.group(1)), _ym(m.group(3), m.group(4)), m.group(2)))
    elif source == "nippon":
        for m in re.finditer(r'["\'](/[^"\']*NIMF-MONTHLY-PORTFOLIO-\d{1,2}-([A-Za-z]{3,9})-(\d{2})\.xlsx?)["\']', html,
                             re.I):  # fmt: skip
            links.append(FileLink(urljoin(page_url, m.group(1)), _ym(m.group(2), m.group(3)), "all"))
    elif source == "dsp":
        for m in re.finditer(r'(https://www\.dspim\.com/media/pages/mandatory-disclosures/portfolio-disclosures/'
                             r'[^"\'\s]*?month_?end[^"\'\s]*?\.zip)', html, re.I):  # fmt: skip
            name = m.group(1).rsplit("/", 1)[-1].lower()
            mw = re.search(_MONTH_WORD, name)
            yr = re.findall(r"(20\d{2})", name)
            if mw and yr:
                links.append(FileLink(m.group(1), _ym(mw.group(1), yr[-1]), "all"))
    seen: dict[tuple[str, str], FileLink] = {}
    for link in links:
        seen.setdefault((link.month, link.label), link)
    return sorted(seen.values(), key=lambda x: (x.month, x.label), reverse=True)


PPFAS_PREFIX = {"parag parikh flexicap fund": "PPFCF", "parag parikh elss tax saver fund": "PPTSF",
                "parag parikh conservative hybrid fund": "PPCHF", "parag parikh arbitrage fund": "PPAF",
                "parag parikh dynamic asset allocation fund": "PPDAAF", "parag parikh liquid fund": "PPLF",
                "parag parikh largecap fund": "PPLCF"}  # fmt: skip


def pick_links(source: str, links: list[FileLink], key: str, months: int) -> list[FileLink]:
    """The files to download for one scheme: its own file (PPFAS) or the house-wide file, latest `months` months."""
    if source == "ppfas":
        prefix = PPFAS_PREFIX.get(key)
        mine = [x for x in links if x.label == prefix] if prefix else []
        chosen = (
            mine or links
        )  # unknown scheme: fall back to every PPFAS file of the month (matched after parsing)
    else:
        chosen = links
    wanted = sorted({x.month for x in chosen}, reverse=True)[: max(1, months)]
    return [x for x in chosen if x.month in wanted]


def fetched_at_now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")
